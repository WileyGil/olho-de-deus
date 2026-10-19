import streamlit as st
import pandas as pd
import xml.etree.ElementTree as ET
from itertools import zip_longest
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
import re
from datetime import datetime
import io

# ==========================================
# 1. FUNÇÕES AUXILIARES
# ==========================================
def limpar_numero(valor):
    if pd.isna(valor):
        return ""
    val_str = str(valor).strip()
    if val_str.endswith('.0'):
        val_str = val_str[:-2]
    return re.sub(r'[^0-9]', '', val_str)

def limpar_ean(valor):
    if pd.isna(valor):
        return ""
    val_str = str(valor).strip().upper()
    if val_str in ('SEM GTIN', 'S/GTIN', 'NONE', 'NAN', ''):
        return ""
    if val_str.endswith('.0'):
        val_str = val_str[:-2]
    
    # Tratamento contra notação científica do Excel ou ERP
    if 'E+' in val_str or 'E-' in val_str:
        try:
            val_limpo = val_str.replace(',', '.')
            val_str = str(int(float(val_limpo)))
        except:
            pass

    apenas_numeros = re.sub(r'[^0-9]', '', val_str)
    return apenas_numeros.lstrip('0')

def converte_float_br(valor):
    if pd.isna(valor) or str(valor).strip() == '':
        return 0.0
    try:
        val_str = str(valor).strip()
        if '.' in val_str and ',' in val_str:
            val_str = val_str.replace('.', '').replace(',', '.')
        else:
            val_str = val_str.replace(',', '.')
        return float(val_str)
    except:
        return 0.0

def obter_icms_xml(det_node, ns):
    pICMS = det_node.find('.//ns:pICMS', ns)
    if pICMS is not None and pICMS.text:
        return float(pICMS.text)
    return 0.0

# ==========================================
# 2. INTERFACE STREAMLIT
# ==========================================
st.set_page_config(page_title="Auditoria Olho de Deus", layout="wide")
st.title("Auditoria de Notas")
st.markdown("Faça o upload do arquivo XML da nota e do arquivo do WinThor (.csv ou .xlsx) para cruzar as informações com base absoluta na ordem do XML.")

col1, col2 = st.columns(2)
with col1:
    arquivo_xml = st.file_uploader("Upload do XML (Fornecedor)", type=["xml"])
with col2:
    arquivo_dados = st.file_uploader("Upload do Relatório WinThor", type=["csv", "xlsx"])

if arquivo_xml and arquivo_dados:
    if st.button("Processar Auditoria"):
        with st.spinner("Analisando com base fixa no XML..."):
            # ==========================================
            # 3. EXTRAÇÃO DO XML (NF - O CENTRO DE TUDO)
            # ==========================================
            ns = {'ns': 'http://www.portalfiscal.inf.br/nfe'}
            tree = ET.parse(arquivo_xml)
            root = tree.getroot()

            numero_nf = "DESCONHECIDO"
            nNF_node = root.find('.//ns:ide/ns:nNF', ns)
            if nNF_node is not None:
                numero_nf = nNF_node.text

            itens_xml = []
            for det in root.findall('.//ns:det', ns):
                prod = det.find('ns:prod', ns)

                try:
                    nItem = int(det.attrib.get('nItem', '0'))
                except:
                    nItem = det.attrib.get('nItem', '')

                desc = prod.find('ns:xProd', ns).text if prod.find('ns:xProd', ns) is not None else ''
                ean_raw = prod.find('ns:cEAN', ns).text if prod.find('ns:cEAN', ns) is not None else ''
                ean_trib_raw = prod.find('ns:cEANTrib', ns).text if prod.find('ns:cEANTrib', ns) is not None else ''

                ean_exibicao = limpar_numero(ean_raw) if limpar_ean(ean_raw) != '' else limpar_numero(ean_trib_raw)
                ean_comp = limpar_ean(ean_raw)
                ean_trib_comp = limpar_ean(ean_trib_raw)

                ncm = limpar_numero(prod.find('ns:NCM', ns).text) if prod.find('ns:NCM', ns) is not None else ''
                qtd = float(prod.find('ns:qCom', ns).text) if prod.find('ns:qCom', ns) is not None else 0.0
                vUn = float(prod.find('ns:vUnCom', ns).text) if prod.find('ns:vUnCom', ns) is not None else 0.0
                vTot = float(prod.find('ns:vProd', ns).text) if prod.find('ns:vProd', ns) is not None else 0.0
                icms = obter_icms_xml(det, ns)

                pRedBC_node = det.find('.//ns:pRedBC', ns)
                pRedBC = float(pRedBC_node.text) if pRedBC_node is not None and pRedBC_node.text else 0.0

                itens_xml.append({
                    'Linha_NF': nItem,
                    'Descricao_NF': desc,
                    'EAN_NF': ean_exibicao,
                    'EAN_Comp_NF': ean_comp,
                    'EANTrib_Comp_NF': ean_trib_comp,
                    'NCM_NF': ncm,
                    'Qtd_NF': qtd,
                    'Vl_Unit_NF': vUn,
                    'Vl_Total_NF': vTot,
                    'ICMS_NF': icms,
                    'pRedBC_NF': pRedBC
                })

            data_simples = datetime.now().strftime('%Y%m%d')
            nome_arquivo_saida = f'relatorio_NF_{numero_nf}_{data_simples}.xlsx'

            # ==========================================
            # 4. EXTRAÇÃO DA PLANILHA (WINTHOR)
            # ==========================================
            if arquivo_dados.name.lower().endswith('.xlsx'):
                df_csv = pd.read_excel(arquivo_dados, dtype=str)
            else:
                df_csv = pd.read_csv(arquivo_dados, sep=';', encoding='latin1', dtype=str)

            df_csv.columns = df_csv.columns.str.strip()

            col_tipo_calc = next((c for c in df_csv.columns if 'tipo' in c.lower() and ('cálculo' in c.lower() or 'calculo' in c.lower())), None)
            if not col_tipo_calc:
                col_tipo_calc = df_csv.columns[-2]

            itens_csv = []
            for index, row in df_csv.iterrows():
                ean_win_raw = row.get('Cód. barra', '')
                itens_csv.append({
                    'Linha_WINTHOR': index + 1,
                    'Cod_Produto_WINTHOR': limpar_numero(row.get('Cód. produto', '')),
                    'EAN_WINTHOR': limpar_numero(ean_win_raw),
                    'EAN_Comp_WINTHOR': limpar_ean(ean_win_raw),
                    'Descricao_WINTHOR': str(row.get('Descrição', '')).strip(),
                    'Qtd_WINTHOR': converte_float_br(row.get('Qtde Pré-Entrada', 0.0)),
                    'NCM_WINTHOR': limpar_numero(row.get('NCM', '')),
                    'ICMS_WINTHOR': converte_float_br(row.get('%ICMS', 0.0)),
                    '% ICMS red. Whintor': converte_float_br(row.get('% ICMS red.', 0.0)),
                    '%Alíq.Ext. ST Guia_Whintor': converte_float_br(row.get('%Alíq.Ext.ST Guia', 0.0)),
                    'Tipo_Calculo_WINTHOR': str(row.get(col_tipo_calc, '')).strip()
                })

            # ==========================================
            # 5. CRUZAMENTO BASEADO ESTRITAMENTE NA ORDEM DO XML
            # ==========================================
            dados_relatorio = []

            for item_x, item_c in zip_longest(itens_xml, itens_csv, fillvalue=None):
                linha = {}

                if item_x:
                    linha.update(item_x)
                else:
                    linha.update({
                        'Linha_NF': '-', 'Descricao_NF': '-', 'EAN_NF': '-', 'NCM_NF': '-',
                        'Qtd_NF': '-', 'Vl_Unit_NF': '-', 'Vl_Total_NF': '-', 'ICMS_NF': '-', 'pRedBC_NF': '-'
                    })

                if item_c:
                    linha.update(item_c)
                else:
                    linha.update({
                        'Linha_WINTHOR': '-', 'Cod_Produto_WINTHOR': '-', 'EAN_WINTHOR': '-', 'Descricao_WINTHOR': '-',
                        'Qtd_WINTHOR': '-', 'NCM_WINTHOR': '-', 'ICMS_WINTHOR': '-', '% ICMS red. Whintor': '-',
                        '%Alíq.Ext. ST Guia_Whintor': '-', 'Tipo_Calculo_WINTHOR': '-'
                    })

                if item_x and item_c:
                    linha['Status_Geral'] = 'PAREADO'

                    ean_x_comp = item_x.get('EAN_Comp_NF', '')
                    eantrib_x_comp = item_x.get('EANTrib_Comp_NF', '')
                    ean_win_comp = item_c.get('EAN_Comp_WINTHOR', '')

                    if ean_win_comp != '' and (ean_win_comp == ean_x_comp or ean_win_comp == eantrib_x_comp):
                        linha['Status_EAN'] = 'OK'
                    elif ean_win_comp == '' and ean_x_comp == '' and eantrib_x_comp == '':
                        linha['Status_EAN'] = 'OK'
                    else:
                        linha['Status_EAN'] = 'DIVERGENTE'

                    icms_xml = round(item_x['ICMS_NF'], 2)
                    icms_csv = round(item_c['ICMS_WINTHOR'], 2)
                    st_guia = round(item_c['%Alíq.Ext. ST Guia_Whintor'], 2)
                    if icms_xml == icms_csv or icms_xml == st_guia:
                        linha['Status_ICMS_Aliq'] = 'OK'
                    else:
                        linha['Status_ICMS_Aliq'] = 'DIVERGENTE'

                    pred_xml = round(item_x['pRedBC_NF'], 4)
                    red_csv = round(item_c['% ICMS red. Whintor'], 4)
                    if pred_xml == red_csv:
                        linha['Status_Red_BC'] = 'OK'
                    else:
                        linha['Status_Red_BC'] = 'DIVERGENTE'

                    tipo_calc = str(item_c.get('Tipo_Calculo_WINTHOR', '')).strip().upper()
                    if tipo_calc == 'N':
                        linha['Tipo_Calculo_ST'] = 'NOTA FISCAL'
                    elif tipo_calc == 'G':
                        linha['Tipo_Calculo_ST'] = 'GUIA'
                    else:
                        linha['Tipo_Calculo_ST'] = item_c['Tipo_Calculo_WINTHOR']

                    linha['Status_NCM'] = 'OK' if str(item_x['NCM_NF']) == str(item_c['NCM_WINTHOR']) else 'DIVERGENTE'
                    linha['Status_Qtd'] = 'OK' if round(item_x['Qtd_NF'], 4) == round(item_c['Qtd_WINTHOR'], 4) else 'DIVERGENTE'

                elif item_x and not item_c:
                    linha['Status_Geral'] = 'SOBRA NF (SEM PAR)'
                    for st_col in ['Status_EAN', 'Status_ICMS_Aliq', 'Status_Red_BC', 'Tipo_Calculo_ST', 'Status_NCM', 'Status_Qtd']:
                        linha[st_col] = '-'

                elif item_c and not item_x:
                    linha['Status_Geral'] = 'SOBRA WINTHOR (SEM PAR)'
                    for st_col in ['Status_EAN', 'Status_ICMS_Aliq', 'Status_Red_BC', 'Tipo_Calculo_ST', 'Status_NCM', 'Status_Qtd']:
                        linha[st_col] = '-'

                dados_relatorio.append(linha)

            # ==========================================
            # 6. CRIANDO EXCEL NA MEMÓRIA (BUFFER)
            # ==========================================
            df_final = pd.DataFrame(dados_relatorio)
            colunas_ordem = [
                'Linha_NF', 'Linha_WINTHOR',
                'Cod_Produto_WINTHOR', 'EAN_NF', 'EAN_WINTHOR', 'Status_EAN',
                'Descricao_NF', 'Descricao_WINTHOR',
                'Qtd_NF', 'Qtd_WINTHOR', 'Status_Qtd',
                'NCM_NF', 'NCM_WINTHOR', 'Status_NCM',
                'ICMS_NF', 'ICMS_WINTHOR', '%Alíq.Ext. ST Guia_Whintor', 'Status_ICMS_Aliq',
                'pRedBC_NF', '% ICMS red. Whintor', 'Status_Red_BC',
                'Tipo_Calculo_ST'
            ]
            df_final = df_final[colunas_ordem]

            excel_buffer = io.BytesIO()
            with pd.ExcelWriter(excel_buffer, engine='openpyxl') as writer:
                df_final.to_excel(writer, sheet_name='Comparativo_ICMS_e_EAN', index=False, startrow=2)
                
                workbook = writer.book
                worksheet = workbook['Comparativo_ICMS_e_EAN']
                
                worksheet.sheet_view.showGridLines = False
                worksheet.sheet_view.zoomScale = 80

                worksheet.freeze_panes = 'A4'
                worksheet.auto_filter.ref = f"A3:V{len(df_final) + 3}"

                fill_zebrado = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
                fill_cab_status_divergente = PatternFill(start_color="9C0006", end_color="9C0006", fill_type="solid")
                fill_cab_status_convergente = PatternFill(start_color="006100", end_color="006100", fill_type="solid")

                font_super_header_xml = Font(bold=True, size=11, color="FFFFFF")
                font_super_header_erp = Font(bold=True, size=11, color="333333")
                font_header_default_xml = Font(bold=True, size=10, color="FFFFFF")
                font_header_default_erp = Font(bold=True, size=10, color="333333")
                font_header_status = Font(bold=True, size=10, color="FFFFFF")

                color_destaque_xml = "1F4E78"
                color_neutro_geral = "E0E2E5"

                fill_cab_xml = PatternFill(start_color=color_destaque_xml, end_color=color_destaque_xml, fill_type="solid")
                fill_cab_neutro = PatternFill(start_color=color_neutro_geral, end_color=color_neutro_geral, fill_type="solid")

                borda_branca = Border(
                    left=Side(style='thin', color='FFFFFF'), right=Side(style='thin', color='FFFFFF'),
                    top=Side(style='thin', color='FFFFFF'), bottom=Side(style='thin', color='FFFFFF')
                )
                borda_divisoria_direita = Border(right=Side(style='medium', color='D9D9D9'))

                super_grupos = [
                    ("IDENTIFICAÇÃO", 1, 2, fill_cab_neutro, font_super_header_erp),
                    ("CÓDIGO E CÓDIGO DE BARRAS (EAN) [XML]", 3, 6, fill_cab_xml, font_super_header_xml),
                    ("DESCRIÇÕES COMPARATIVAS", 7, 8, fill_cab_neutro, font_super_header_erp),
                    ("QUANTIDADES [XML]", 9, 11, fill_cab_xml, font_super_header_xml),
                    ("CLASSIFICAÇÃO FISCAL (NCM) [XML]", 12, 14, fill_cab_xml, font_super_header_xml),
                    ("ALÍQUOTAS DE ICMS [XML]", 15, 18, fill_cab_xml, font_super_header_xml),
                    ("REDUÇÃO DE BASE DE CÁLCULO [XML]", 19, 21, fill_cab_xml, font_super_header_xml),
                    ("TIPO CÁLCULO ST", 22, 22, fill_cab_neutro, font_super_header_erp)
                ]

                for label, col_start, col_end, fill_obj, font_obj in super_grupos:
                    worksheet.merge_cells(start_row=1, start_column=col_start, end_row=2, end_column=col_end)
                    cell = worksheet.cell(row=1, column=col_start)
                    cell.value = label
                    cell.font = font_obj
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

                    for r in (1, 2):
                        for c in range(col_start, col_end + 1):
                            cell_target = worksheet.cell(row=r, column=c)
                            cell_target.fill = fill_obj
                            cell_target.border = borda_branca
                            if c == col_end and col_end != 22:
                                cell_target.border = borda_divisoria_direita

                worksheet.row_dimensions[1].height = 18
                worksheet.row_dimensions[2].height = 18
                worksheet.row_dimensions[3].height = 25

                status_columns = ['Status_EAN', 'Status_Qtd', 'Status_NCM', 'Status_ICMS_Aliq', 'Status_Red_BC']
                status_has_divergence = {}
                for sc in status_columns:
                    status_has_divergence[sc] = (df_final[sc] == 'DIVERGENTE').any() if sc in df_final.columns else False

                headers = colunas_ordem
                for col_idx, col_name in enumerate(headers, 1):
                    cell = worksheet.cell(row=3, column=col_idx)
                    cell.value = col_name
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

                    default_border = borda_divisoria_direita if col_idx in (2, 6, 8, 11, 14, 18, 21, 22) else borda_branca
                    cell.border = default_border

                    if col_name in status_has_divergence:
                        if status_has_divergence[col_name]:
                            cell.fill = fill_cab_status_divergente
                            cell.font = font_header_status
                        else:
                            cell.fill = fill_cab_status_convergente
                            cell.font = font_header_status
                    elif 'NF' in col_name or col_name in ('EAN_NF', 'Qtd_NF', 'NCM_NF', 'ICMS_NF', 'pRedBC_NF'):
                        cell.fill = fill_cab_xml
                        cell.font = font_header_default_xml
                    else:
                        cell.fill = fill_cab_neutro
                        cell.font = font_header_default_erp

                fill_cinza = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
                fill_vermelho_claro = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                fill_verde_oliva_claro = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

                font_divergente = Font(color="9C0006", bold=True)
                font_normal = Font(color="000000", bold=False)

                align_center_cols = [
                    'Linha_NF', 'Linha_WINTHOR', 'EAN_NF', 'EAN_WINTHOR', 'Qtd_NF', 'Qtd_WINTHOR',
                    'NCM_NF', 'NCM_WINTHOR', 'ICMS_NF', 'ICMS_WINTHOR', '%Alíq.Ext. ST Guia_Whintor',
                    'pRedBC_NF', '% ICMS red. Whintor'
                ]

                status_dependents = {
                    'Status_EAN': ['EAN_NF', 'EAN_WINTHOR'],
                    'Status_Qtd': ['Qtd_NF', 'Qtd_WINTHOR'],
                    'Status_NCM': ['NCM_NF', 'NCM_WINTHOR'],
                    'Status_ICMS_Aliq': ['ICMS_NF', 'ICMS_WINTHOR', '%Alíq.Ext. ST Guia_Whintor'],
                    'Status_Red_BC': ['pRedBC_NF', '% ICMS red. Whintor']
                }

                for row in range(4, len(df_final) + 4):
                    linha_dict = dados_relatorio[row-4]
                    is_even = (row % 2 == 0)

                    for col_idx in range(1, len(headers) + 1):
                        cell = worksheet.cell(row=row, column=col_idx)

                        if is_even:
                            cell.fill = fill_zebrado

                        default_border = borda_divisoria_direita if col_idx in (2, 6, 8, 11, 14, 18, 21, 22) else borda_branca
                        cell.border = default_border
                        col_name = headers[col_idx - 1]

                        current_align = cell.alignment or Alignment()
                        cell.alignment = Alignment(horizontal=current_align.horizontal or 'left', vertical='center', wrap_text=current_align.wrap_text)

                        if col_name == 'Cod_Produto_WINTHOR' and str(cell.value).strip() not in ('', '-'):
                            try:
                                cell.value = int(float(cell.value))
                                cell.number_format = '0'
                            except ValueError:
                                pass
                        elif col_name in ('NCM_NF', 'NCM_WINTHOR') and str(cell.value).strip() not in ('', '-'):
                            try:
                                cell.value = int(float(cell.value))
                                cell.number_format = '00000000'
                            except ValueError:
                                pass
                        elif col_name in ('EAN_NF', 'EAN_WINTHOR') and str(cell.value).strip() not in ('', '-'):
                            cell.value = str(cell.value).strip()
                            cell.number_format = '@'
                        elif isinstance(cell.value, (int, float)) and cell.value != '-':
                            if 'pRedBC' in col_name or 'red. Whintor' in col_name:
                                cell.number_format = '#,##0.0000'
                            elif 'Qtd' in col_name or 'ICMS' in col_name or 'Whintor' in col_name:
                                cell.number_format = '#,##0.00'

                        if col_name == 'Cod_Produto_WINTHOR':
                            cell.alignment = Alignment(horizontal='right', vertical='center')
                        elif col_name in align_center_cols or 'Status' in col_name:
                            cell.alignment = Alignment(horizontal='center', vertical='center')

                    if 'SOBRA' in str(linha_dict.get('Status_Geral', '')):
                        for col_idx in range(1, len(headers) + 1):
                            worksheet.cell(row=row, column=col_idx).fill = fill_cinza
                        continue

                    for col_status, colunas_dados in status_dependents.items():
                        if col_status in headers:
                            idx_status = headers.index(col_status) + 1
                            cell_status = worksheet.cell(row=row, column=idx_status)
                            val_status = cell_status.value

                            if val_status == 'DIVERGENTE':
                                cell_status.fill = fill_vermelho_claro
                                for col_afetada in colunas_dados:
                                    if col_afetada in headers:
                                        cell_dado = worksheet.cell(row=row, column=headers.index(col_afetada) + 1)
                                        cell_dado.font = font_divergente
                            elif val_status == 'OK':
                                cell_status.fill = fill_verde_oliva_claro
                                for col_afetada in colunas_dados:
                                    if col_afetada in headers:
                                        cell_dado = worksheet.cell(row=row, column=headers.index(col_afetada) + 1)
                                        cell_dado.font = font_normal

                for col in worksheet.columns:
                    max_length = 0
                    column_letter = get_column_letter(col[0].column)
                    for cell in col:
                        try:
                            if cell.value is not None:
                                if len(str(cell.value)) > max_length:
                                    max_length = len(str(cell.value))
                        except:
                            pass
                    worksheet.column_dimensions[column_letter].width = max(max_length + 4, 12)

                # ==========================================
                # 7. BLOQUEIO DA PLANILHA (PERMITIR FILTROS, ORDENAÇÃO E DESBLOQUEIO DE CÓDIGOS)
                # ==========================================
                worksheet.protection.sheet = True
                worksheet.protection.sort = False          # Libera ordenação de A a Z / Z a A
                worksheet.protection.autoFilter = False      # Libera os filtros suspensos nas colunas
                worksheet.protection.selectLockedCells = True
                worksheet.protection.selectUnlockedCells = True

                # Colunas a serem desbloqueadas para edição livre: Cod_Produto_WINTHOR, EAN_NF, EAN_WINTHOR
                colunas_para_desbloquear = ['Cod_Produto_WINTHOR', 'EAN_NF', 'EAN_WINTHOR']
                indices_desbloquear = [colunas_ordem.index(c) + 1 for c in colunas_para_desbloquear if c in colunas_ordem]

                for row in range(4, len(df_final) + 4):
                    for col_idx in range(1, len(colunas_ordem) + 1):
                        cell = worksheet.cell(row=row, column=col_idx)
                        if col_idx in indices_desbloquear:
                            cell.protection = Protection(locked=False)
                        else:
                            cell.protection = Protection(locked=True)

            excel_buffer.seek(0)
            
            st.success("Auditoria finalizada com sucesso!")
            st.download_button(
                label="📥 Baixar Relatório Executivo (Excel)",
                data=excel_buffer.getvalue(),
                file_name=nome_arquivo_saida,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
